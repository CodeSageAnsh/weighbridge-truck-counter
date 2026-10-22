"""
Count trucks crossing a weighbridge in CCTV footage.

    python -m truckcount.cli --video bridge.mp4 --line 596 134 362 --out out/

Outputs
  out/counted.mp4         annotated video (line, boxes, running total)
  out/shots/truckN_*.jpg  clearest full frame per truck (CCTV clock readable)
  out/truck_count.xlsx    proof sheet: #, video time, blank real-time column, screenshot
"""
import argparse
import os

import cv2
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .line_counter import Line, LineCounter

COCO_VEHICLES = {2: "car", 5: "bus", 7: "truck"}


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--video", required=True)
    p.add_argument("--model", default="yolo11n.pt", help="YOLO weights (COCO or custom)")
    p.add_argument("--line", nargs=3, type=int, metavar=("X", "TOP", "BOTTOM"), required=True,
                   help="vertical count line, clipped to the bridge slab")
    p.add_argument("--out", default="output")
    p.add_argument("--conf", type=float, default=0.45)
    p.add_argument("--stride", type=int, default=2, help="process every Nth frame")
    p.add_argument("--enter-sec", type=float, default=1.0, help="line must be covered this long to count")
    p.add_argument("--clear-sec", type=float, default=2.5, help="line must be clear this long to re-arm")
    p.add_argument("--classes", type=int, nargs="*", default=list(COCO_VEHICLES),
                   help="class ids to count (default: COCO car/bus/truck)")
    return p.parse_args()


def video_meta(path):
    cap = cv2.VideoCapture(path)
    fps = cap.get(cv2.CAP_PROP_FPS) or 25
    w, h = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH)), int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
    cap.release()
    if w == 0:
        raise SystemExit(f"Cannot open video: {path}")
    return fps, w, h


def annotate(frame, line, boxes, total, state):
    out = frame.copy()
    col = (0, 0, 255) if state == "COVERED" else (0, 200, 255)
    cv2.line(out, (line.x, line.top), (line.x, line.bottom), col, 4)
    for x1, y1, x2, y2, on_line in boxes:
        cv2.rectangle(out, (x1, y1), (x2, y2), (0, 255, 0) if on_line else (170, 170, 170), 2)
    cv2.rectangle(out, (0, 0), (out.shape[1], 46), (0, 0, 0), -1)
    tag = "TRUCK ON LINE" if state == "COVERED" else "clear"
    cv2.putText(out, f"TRUCKS COUNTED: {total}   [{tag}]", (10, 32),
                cv2.FONT_HERSHEY_SIMPLEX, 0.9, (0, 255, 255), 2)
    return out


def write_proof_sheet(rows, total, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Truck count"
    navy = "1F3864"
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("A1:E1")
    ws["A1"] = f"WEIGHBRIDGE TRUCK COUNT   (TOTAL {total})"
    ws["A1"].font = Font(bold=True, size=13, color=navy)
    heads = ["#", "Video time", "Real clock time\n(read from screenshot)", "Confidence", "Screenshot"]
    for i, (h, w) in enumerate(zip(heads, [8, 12, 22, 12, 60]), 1):
        c = ws.cell(3, i, h)
        c.font, c.alignment, c.border = Font(bold=True, color="FFFFFF"), center, border
        c.fill = PatternFill("solid", fgColor=navy)
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, start=4):
        ws.row_dimensions[r].height = 200
        for col, val in enumerate([row["n"], row["ts"], "", row["conf"]], 1):
            ws.cell(r, col, val).alignment = center
        if os.path.exists(row["shot"]):
            img = XLImage(row["shot"])
            img.width, img.height = 420, 236
            ws.add_image(img, f"E{r}")
        for col in range(1, 6):
            ws.cell(r, col).border = border
    wb.save(path)


def main():
    a = parse_args()
    from ultralytics import YOLO  # imported late so --help works without torch

    shots = os.path.join(a.out, "shots")
    os.makedirs(shots, exist_ok=True)
    fps, w, h = video_meta(a.video)
    eff_fps = fps / a.stride
    line = Line(*a.line)
    counter = LineCounter(enter_frames=max(1, int(a.enter_sec * eff_fps)),
                          clear_frames=max(1, int(a.clear_sec * eff_fps)))
    writer = cv2.VideoWriter(os.path.join(a.out, "counted.mp4"),
                             cv2.VideoWriter_fourcc(*"mp4v"), eff_fps, (w, h))

    rows, best = [], None  # best = (conf, clean frame, timestamp) for the truck on the line
    results = YOLO(a.model).predict(source=a.video, classes=a.classes, conf=a.conf,
                                    vid_stride=a.stride, stream=True, verbose=False)
    for i, res in enumerate(results):
        frame = res.orig_img
        sec = i * a.stride / fps
        ts = f"{int(sec // 60):02d}:{int(sec % 60):02d}"

        boxes, covered, frame_conf = [], False, 0.0
        if res.boxes is not None and len(res.boxes):
            for cf, box in zip(res.boxes.conf.cpu().numpy(), res.boxes.xyxy.cpu().numpy()):
                x1, y1, x2, y2 = map(int, box)
                on_line = line.covered_by(x1, y1, x2, y2)
                boxes.append((x1, y1, x2, y2, on_line))
                if on_line:
                    covered, frame_conf = True, max(frame_conf, float(cf))

        if covered and frame_conf > (best[0] if best else 0):
            best = (frame_conf, frame.copy(), ts)

        if counter.update(covered):
            conf, shot_frame, shot_ts = best or (0.0, frame, ts)
            shot = os.path.join(shots, f"truck{counter.total}_{shot_ts.replace(':', '')}.jpg")
            cv2.imwrite(shot, shot_frame)
            rows.append({"n": counter.total, "ts": shot_ts, "conf": round(conf, 2), "shot": shot})
            print(f"  TRUCK #{counter.total} @ {shot_ts}  conf={conf:.2f}")
        elif counter.state == "CLEAR":
            best = None

        writer.write(annotate(frame, line, boxes, counter.total, counter.state))

    writer.release()
    write_proof_sheet(rows, counter.total, os.path.join(a.out, "truck_count.xlsx"))
    print(f"\nTOTAL TRUCKS: {counter.total}  ->  {a.out}")


if __name__ == "__main__":
    main()
